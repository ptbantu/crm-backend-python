"""
文档模板服务
"""
from typing import Optional, List
from sqlalchemy.ext.asyncio import AsyncSession
import uuid
from io import BytesIO

from common.models.doc_template import DocTemplate, TemplateTypeEnum
from foundation_service.repositories.doc_template_repository import DocTemplateRepository
from foundation_service.schemas.document_template import (
    DocTemplateCreateRequest,
    DocTemplateUpdateRequest,
    DocTemplateResponse,
    DocTemplateListResponse,
)
from common.utils.logger import get_logger
from common.exceptions import BusinessException

logger = get_logger(__name__)


class DocumentTemplateService:
    """文档模板服务"""

    def __init__(self, db: AsyncSession):
        self.db = db
        self.template_repo = DocTemplateRepository(db)

    async def create_template(
        self,
        request: DocTemplateCreateRequest,
        created_by: Optional[str] = None
    ) -> DocTemplateResponse:
        """创建文档模板"""
        # 检查模板代码是否已存在
        existing = await self.template_repo.get_by_code(request.template_code)
        if existing:
            raise BusinessException(detail="模板代码已存在", status_code=400)

        # 如果没有提供占位符列表，尝试从文件解析
        placeholders = request.placeholders
        if not placeholders and request.file_oss_key.endswith('.docx'):
            try:
                placeholders = await self._parse_placeholders_from_docx(request.file_oss_key)
            except Exception as e:
                logger.warning(f"解析模板占位符失败: {str(e)}")
                placeholders = []

        # 创建模板
        template = DocTemplate(
            id=str(uuid.uuid4()),
            template_code=request.template_code,
            template_name=request.template_name,
            template_type=request.template_type,
            entity_id=request.entity_id,
            language=request.language,
            file_oss_key=request.file_oss_key,
            file_name=request.file_name,
            file_size_kb=request.file_size_kb,
            placeholders=placeholders,
            description=request.description,
            is_active=True,
            version=1,
            created_by=created_by,
            updated_by=created_by
        )

        await self.template_repo.create(template)
        await self.db.commit()
        await self.db.refresh(template)

        logger.info(f"创建文档模板成功: {template.template_code}")
        return DocTemplateResponse.model_validate(template)

    async def get_template(self, template_id: str) -> DocTemplateResponse:
        """获取模板详情"""
        template = await self.template_repo.get_by_id(template_id)
        if not template:
            raise BusinessException(detail="模板不存在", status_code=404)

        return DocTemplateResponse.model_validate(template)

    async def update_template(
        self,
        template_id: str,
        request: DocTemplateUpdateRequest,
        updated_by: Optional[str] = None
    ) -> DocTemplateResponse:
        """更新模板"""
        template = await self.template_repo.get_by_id(template_id)
        if not template:
            raise BusinessException(detail="模板不存在", status_code=404)

        # 更新字段
        if request.template_name is not None:
            template.template_name = request.template_name
        if request.description is not None:
            template.description = request.description
        if request.is_active is not None:
            template.is_active = request.is_active

        template.updated_by = updated_by

        await self.template_repo.update(template)
        await self.db.commit()
        await self.db.refresh(template)

        logger.info(f"更新文档模板成功: {template.template_code}")
        return DocTemplateResponse.model_validate(template)

    async def activate_template(
        self,
        template_id: str,
        updated_by: Optional[str] = None
    ) -> DocTemplateResponse:
        """激活模板"""
        template = await self.template_repo.get_by_id(template_id)
        if not template:
            raise BusinessException(detail="模板不存在", status_code=404)

        template.is_active = True
        template.updated_by = updated_by

        await self.template_repo.update(template)
        await self.db.commit()
        await self.db.refresh(template)

        logger.info(f"激活文档模板: {template.template_code}")
        return DocTemplateResponse.model_validate(template)

    async def deactivate_template(
        self,
        template_id: str,
        updated_by: Optional[str] = None
    ) -> DocTemplateResponse:
        """停用模板"""
        template = await self.template_repo.get_by_id(template_id)
        if not template:
            raise BusinessException(detail="模板不存在", status_code=404)

        template.is_active = False
        template.updated_by = updated_by

        await self.template_repo.update(template)
        await self.db.commit()
        await self.db.refresh(template)

        logger.info(f"停用文档模板: {template.template_code}")
        return DocTemplateResponse.model_validate(template)

    async def list_templates(
        self,
        template_type: Optional[TemplateTypeEnum] = None,
        entity_id: Optional[str] = None,
        is_active: Optional[bool] = None,
        page: int = 1,
        size: int = 20
    ) -> DocTemplateListResponse:
        """查询模板列表"""
        templates, total = await self.template_repo.list_templates(
            template_type=template_type,
            entity_id=entity_id,
            is_active=is_active,
            page=page,
            size=size
        )

        items = [DocTemplateResponse.model_validate(t) for t in templates]
        return DocTemplateListResponse(
            items=items,
            total=total,
            page=page,
            size=size
        )

    async def _parse_placeholders_from_docx(self, oss_key: str) -> List[str]:
        """从Word文档解析占位符（${variable}格式）"""
        try:
            from docx import Document
            import re
            import tempfile
            import os

            # 从 OSS 下载文件
            from foundation_service.services.document_storage_service import DocumentStorageService
            storage_service = DocumentStorageService(self.db)
            file_content = await storage_service.download_from_oss(oss_key)

            # 保存到临时文件
            with tempfile.NamedTemporaryFile(suffix='.docx', delete=False) as temp_file:
                temp_file.write(file_content)
                temp_path = temp_file.name

            try:
                # 打开 Word 文档
                doc = Document(temp_path)

                # 用于存储找到的占位符
                placeholders = set()

                # 正则表达式匹配 ${variable} 格式
                pattern = re.compile(r'\$\{([^}]+)\}')

                # 从段落中提取占位符
                for paragraph in doc.paragraphs:
                    matches = pattern.findall(paragraph.text)
                    placeholders.update(matches)

                # 从表格中提取占位符
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                matches = pattern.findall(paragraph.text)
                                placeholders.update(matches)

                logger.info(f"从模板解析到 {len(placeholders)} 个占位符: {list(placeholders)}")
                return sorted(list(placeholders))

            finally:
                # 清理临时文件
                if os.path.exists(temp_path):
                    os.remove(temp_path)

        except Exception as e:
            logger.error(f"解析模板占位符失败: {str(e)}")
            return []
